"""Unit and integration tests for ResumeParserService and CandidateService."""

import io
import json
from pathlib import Path
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import docx
import pypdf

from database.connection import get_session
from database.models import Base, Candidate, ResumeDocument
from services.candidate_service import CandidateService
from services.resume_parser import COMMON_SKILLS, ResumeParserService


SAMPLE_RESUME_TEXT = """Johnathan Doe
Senior Full-Stack Software Engineer
john.doe@example.com | +1 (555) 234-5678 | San Francisco, CA

PROFESSIONAL SUMMARY
Results-driven software engineer with 8+ years of experience designing and scaling distributed systems,
cloud-native microservices, and modern web applications. Passionate about clean code, performance optimization,
and engineering excellence.

TECHNICAL SKILLS
Languages: Python, TypeScript, JavaScript, SQL, Go, HTML5, CSS3
Frameworks: React, Next.js, FastAPI, Django, Node.js, Express.js
Databases: PostgreSQL, Redis, MongoDB, Elasticsearch
Cloud & DevOps: AWS, Docker, Kubernetes, Terraform, GitHub Actions, CI/CD
Methodologies: Agile, Scrum, Unit Testing, TDD, Microservices

WORK EXPERIENCE
Senior Software Engineer at CloudScale Technologies
Jan 2021 - Present
- Architected and deployed microservices using FastAPI and PostgreSQL handling 20M+ daily requests.
- Integrated Redis caching and Kafka message streaming to reduce API latency by 45%.
- Automated infrastructure deployment using Terraform and GitHub Actions on AWS.

Software Engineer at Apex Innovations
Jun 2017 - Dec 2020
- Built full-stack web applications with React, Next.js, and Django REST Framework.
- Collaborated with product teams in an Agile environment with bi-weekly sprints.

EDUCATION
Bachelor of Technology in Computer Science
University of California, Berkeley
2013 - 2017

CERTIFICATIONS
- AWS Certified Solutions Architect - Associate
- Certified Kubernetes Administrator (CKA)
"""


@pytest.fixture
def parser() -> ResumeParserService:
    return ResumeParserService()


@pytest.fixture
def candidate_service(parser: ResumeParserService) -> CandidateService:
    return CandidateService(parser_service=parser)


# ==============================================================
# ResumeParserService Unit Tests
# ==============================================================

def test_common_skills_vocabulary_size():
    """Verify COMMON_SKILLS has at least 300 technical skills."""
    assert len(COMMON_SKILLS) >= 300
    assert "Python" in COMMON_SKILLS
    assert "FastAPI" in COMMON_SKILLS
    assert "Docker" in COMMON_SKILLS
    assert "PostgreSQL" in COMMON_SKILLS
    assert "Kubernetes" in COMMON_SKILLS
    assert "AWS" in COMMON_SKILLS


def test_compute_file_hash(parser: ResumeParserService):
    """Verify SHA-256 deterministic computation."""
    data = b"TalentMatch AI Resume Ingestion Test Data"
    h1 = parser.compute_file_hash(data)
    h2 = parser.compute_file_hash(data)
    assert len(h1) == 64
    assert h1 == h2
    assert h1 != parser.compute_file_hash(b"different data")


def test_validate_file_valid_types(parser: ResumeParserService):
    """Test validation of valid PDF, DOCX, and TXT files."""
    # Plain text
    txt_res = parser.validate_file(b"Sample resume text", "resume.txt")
    assert txt_res["valid"] is True
    assert txt_res["file_type"] == "txt"
    assert txt_res["error"] is None

    # PDF header
    pdf_res = parser.validate_file(b"%PDF-1.4 test binary data", "resume.pdf")
    assert pdf_res["valid"] is True
    assert pdf_res["file_type"] == "pdf"

    # DOCX header (ZIP PK\x03\x04)
    docx_res = parser.validate_file(b"PK\x03\x04test docx binary", "resume.docx")
    assert docx_res["valid"] is True
    assert docx_res["file_type"] == "docx"


def test_validate_file_edge_cases(parser: ResumeParserService):
    """Test empty files, oversize files, unsupported extensions, and invalid headers."""
    # Empty file
    res_empty = parser.validate_file(b"", "resume.txt")
    assert res_empty["valid"] is False
    assert "empty" in res_empty["error"].lower()

    # Oversize file
    large_bytes = b"0" * (11 * 1024 * 1024)
    res_large = parser.validate_file(large_bytes, "resume.txt", max_size_mb=10)
    assert res_large["valid"] is False
    assert "exceeds" in res_large["error"].lower()

    # Unsupported extension
    res_ext = parser.validate_file(b"test", "resume.exe")
    assert res_ext["valid"] is False
    assert "unsupported" in res_ext["error"].lower()

    # Invalid PDF header
    res_corrupt_pdf = parser.validate_file(b"NotAPdfFileHeader", "corrupt.pdf")
    assert res_corrupt_pdf["valid"] is False
    assert "%pdf-" in res_corrupt_pdf["error"].lower()


def test_parse_txt(parser: ResumeParserService):
    """Test text parsing with UTF-8 and Latin-1."""
    text_data = "John Doe\nPython Developer\njohn@example.com"
    res = parser.parse_txt(text_data.encode("utf-8"), "test.txt")
    assert res["text"] == text_data
    assert len(res["warnings"]) == 0

    # Latin-1 characters
    latin_bytes = "René Descartes\nrene@example.com".encode("latin-1")
    res_latin = parser.parse_txt(latin_bytes, "latin.txt")
    assert "Ren" in res_latin["text"]


def test_parse_docx_in_memory(parser: ResumeParserService):
    """Test DOCX document parsing with paragraphs and tables."""
    doc = docx.Document()
    doc.add_paragraph("Alice Johnson")
    doc.add_paragraph("alice@example.com | (123) 456-7890")
    doc.add_paragraph("Skills: Python, React, PostgreSQL")

    # Add table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Backend Engineer"
    table.cell(1, 0).text = "Company"
    table.cell(1, 1).text = "Acme Corp"

    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    res = parser.parse_docx(docx_bytes, "test.docx")
    assert "Alice Johnson" in res["text"]
    assert "Backend Engineer" in res["text"]
    assert "Acme Corp" in res["text"]
    assert len(res["warnings"]) == 0


def test_parse_pdf_in_memory_and_scanned_detection(parser: ResumeParserService):
    """Test PDF text extraction and scanned PDF detection heuristic."""
    # Create valid PDF in memory using pypdf Writer
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    pdf_io = io.BytesIO()
    writer.write(pdf_io)
    blank_pdf_bytes = pdf_io.getvalue()

    res = parser.parse_pdf(blank_pdf_bytes, "scanned.pdf")
    assert res["page_count"] == 1
    # Because blank page has 0 text chars, scanned detection should trigger
    assert res["is_scanned"] is True
    assert len(res["warnings"]) >= 1


def test_extract_candidate_info(parser: ResumeParserService):
    """Test structured entity extraction from a complete resume."""
    extracted = parser.extract_candidate_info(SAMPLE_RESUME_TEXT)

    assert extracted["name"] == "Johnathan Doe"
    assert extracted["email"] == "john.doe@example.com"
    assert "555" in (extracted["phone"] or "")

    # Skills
    skills = extracted["skills"]
    assert "Python" in skills
    assert "TypeScript" in skills
    assert "FastAPI" in skills
    assert "Docker" in skills
    assert "PostgreSQL" in skills
    assert "Kubernetes" in skills

    # Summary
    assert len(extracted["summary"]) > 0
    assert "Results-driven" in extracted["summary"]

    # Experience
    assert len(extracted["experience"]) >= 1
    exp0 = extracted["experience"][0]
    assert "Senior Software Engineer" in exp0["title"]
    assert "CloudScale Technologies" in exp0["company"]

    # Education
    assert len(extracted["education"]) >= 1
    edu0 = extracted["education"][0]
    assert "Technology" in edu0["degree"] or "Bachelor" in edu0["degree"]
    assert "Berkeley" in edu0["institution"]

    # Certifications
    assert len(extracted["certifications"]) >= 1
    assert any("AWS" in c for c in extracted["certifications"])


# ==============================================================
# CandidateService Integration Tests
# ==============================================================

def test_process_uploaded_file_lifecycle(candidate_service: CandidateService):
    """Test the full ingestion pipeline: validate -> extract -> save -> query -> delete."""
    file_bytes = SAMPLE_RESUME_TEXT.encode("utf-8")
    filename = "johnathan_doe_resume.txt"

    # Ingest candidate
    result = candidate_service.process_uploaded_file(file_bytes, filename)
    assert result["success"] is True
    assert result["status"] == "new"
    cand_id = result["candidate_id"]
    assert cand_id is not None
    assert result["candidate_name"] == "Johnathan Doe"

    # Deduplication test: re-uploading identical file must be detected as duplicate
    dup_res = candidate_service.process_uploaded_file(file_bytes, filename)
    assert dup_res["success"] is False
    assert dup_res["status"] == "duplicate"
    assert dup_res["candidate_id"] == cand_id

    # Retrieve candidate
    cand_details = candidate_service.get_candidate(cand_id)
    assert cand_details is not None
    assert cand_details["name"] == "Johnathan Doe"
    assert cand_details["email"] == "john.doe@example.com"
    assert "Python" in cand_details["skills"]
    assert len(cand_details["resumes"]) == 1
    assert cand_details["resumes"][0]["original_filename"] == filename

    # Update candidate profile
    updated = candidate_service.update_candidate(
        cand_id,
        {
            "name": "Johnathan Doe, Lead Architect",
            "skills": ["Python", "Go", "Rust", "Architecture"],
            "parsing_status": "parsed",
        },
    )
    assert updated is not None
    assert updated["name"] == "Johnathan Doe, Lead Architect"
    assert "Rust" in updated["skills"]

    # Search candidates
    search_results, total = candidate_service.get_candidates(search="Lead Architect")
    assert total >= 1
    assert any(c["id"] == cand_id for c in search_results)

    # Delete candidate
    deleted = candidate_service.delete_candidate(cand_id)
    assert deleted is True

    # Ensure candidate no longer exists
    assert candidate_service.get_candidate(cand_id) is None


def test_process_pasted_text_multiple(candidate_service: CandidateService):
    """Test multi-candidate pasted text ingestion separated by delimiter."""
    pasted = """Alice Walker
Senior React Developer
alice.walker@example.com | +1 555-111-2222
Skills: React, TypeScript, Next.js, Redux, CSS3

---

Bob Vance
DevOps Architect
bob.vance@example.com | +1 555-333-4444
Skills: Kubernetes, Docker, Terraform, AWS, Ansible
"""
    res = candidate_service.process_pasted_text(pasted)
    assert res["total"] == 2
    assert res["success_count"] == 2
    assert res["duplicate_count"] == 0
    assert len(res["results"]) == 2

    # Cleanup
    for item in res["results"]:
        if item.get("candidate_id"):
            candidate_service.delete_candidate(item["candidate_id"])
